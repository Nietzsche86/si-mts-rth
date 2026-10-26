import io
from typing import List, Dict, Any, Optional
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, status
from fastapi.responses import StreamingResponse
from sqlmodel import Session, select
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.api.deps import get_current_user
from app.db.engine import get_session
from app.models import (
    User, Rombel, KeanggotaanRombel,
    Siswa, SiswaDetail, SiswaOrtu, Alamat, SiswaSensitif
)
from app.core.security import encrypt_sensitive_field, decrypt_sensitive_field
from app.core.config import settings

router = APIRouter(prefix="/emis", tags=["EMIS Import & Export"])

EMIS_HEADERS = [
    "NO", "NIK", "NO KK", "NIS", "NISN", "NAMA LENGKAP", "JK",
    "TEMPAT LAHIR", "TANGGAL LAHIR (YYYY-MM-DD)", "NAMA IBU", "NAMA AYAH",
    "ALAMAT JALAN", "DESA/KELURAHAN", "KECAMATAN", "KABUPATEN/KOTA", "PROVINSI"
]

@router.get("/template")
def download_emis_template(
    current_user: User = Depends(get_current_user),
):
    """Mengunduh berkas template Excel EMIS kosong untuk pengisian data siswa baru."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Format EMIS Siswa"

    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0F6E56", end_color="0F6E56", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")

    ws.append(EMIS_HEADERS)
    for col_idx in range(1, len(EMIS_HEADERS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    # Baris contoh
    ws.append([
        1, "3301011205120001", "3301010101100002", "1001", "0081234567",
        "Ahmad Fauzi", "L", "Semarang", "2012-05-12", "Siti Aminah", "Abdul Somad",
        "Jl. Pesantren No. 12", "Kauman", "Semarang Tengah", "Kota Semarang", "Jawa Tengah"
    ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    headers = {"Content-Disposition": 'attachment; filename="template_emis_siswa.xlsx"'}
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)

@router.get("/export/{rombel_id}")
def export_rombel_emis(
    rombel_id: str,
    db: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
):
    """Mengekspor data siswa aktif dalam rombel ke format Excel EMIS Kemenag."""
    rombel = db.get(Rombel, rombel_id)
    if not rombel:
        raise HTTPException(status_code=404, detail="Rombel tidak ditemukan")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"EMIS_{rombel.nama}"

    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0F6E56", end_color="0F6E56", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")

    ws.append(EMIS_HEADERS)
    for col_idx in range(1, len(EMIS_HEADERS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    # Ambil siswa aktif di rombel
    stmt = (
        select(Siswa, KeanggotaanRombel)
        .join(KeanggotaanRombel, KeanggotaanRombel.siswa_id == Siswa.id)
        .where(
            KeanggotaanRombel.rombel_id == rombel_id,
            KeanggotaanRombel.tanggal_keluar == None,
            KeanggotaanRombel.deleted == 0,
            Siswa.deleted == 0,
        )
        .order_by(Siswa.nama_lengkap)
    )
    rows = db.exec(stmt).all()

    for idx, (siswa, keanggotaan) in enumerate(rows, start=1):
        detail = db.exec(select(SiswaDetail).where(SiswaDetail.siswa_id == siswa.id, SiswaDetail.deleted == 0)).first()
        ortu_ayah = db.exec(select(SiswaOrtu).where(SiswaOrtu.siswa_id == siswa.id, SiswaOrtu.hubungan == "ayah", SiswaOrtu.deleted == 0)).first()
        ortu_ibu = db.exec(select(SiswaOrtu).where(SiswaOrtu.siswa_id == siswa.id, SiswaOrtu.hubungan == "ibu", SiswaOrtu.deleted == 0)).first()
        alamat = db.exec(select(Alamat).where(Alamat.pemilik_id == siswa.id, Alamat.deleted == 0)).first()
        sensitif = db.exec(select(SiswaSensitif).where(SiswaSensitif.siswa_id == siswa.id, SiswaSensitif.deleted == 0)).first()

        nik_plain = ""
        kk_plain = ""
        if sensitif:
            if sensitif.nik_siswa_enc:
                nik_plain = decrypt_sensitive_field(sensitif.nik_siswa_enc) or "ERR_DECRYPT"
            if sensitif.no_kk_enc:
                kk_plain = decrypt_sensitive_field(sensitif.no_kk_enc) or "ERR_DECRYPT"

        ws.append([
            idx,
            nik_plain,
            kk_plain,
            siswa.nis,
            siswa.nisn or "",
            siswa.nama_lengkap,
            siswa.jk,
            detail.tempat_lahir if detail else "",
            detail.tanggal_lahir if detail else "",
            ortu_ibu.nama_lengkap if ortu_ibu else "",
            ortu_ayah.nama_lengkap if ortu_ayah else "",
            alamat.jalan if alamat else "",
            alamat.desa_nama if alamat else "",
            alamat.kecamatan_nama if alamat else "",
            alamat.kabupaten_nama if alamat else "",
            alamat.provinsi_nama if alamat else "",
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"emis_rombel_{rombel.nama}.xlsx"
    headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)

@router.post("/import-preview")
async def import_preview(
    file: UploadFile = File(...),
    db: Session = Depends(get_session),
    current_user: User = Depends(get_current_user),
):
    """Membaca dan memvalidasi berkas Excel EMIS sebelum dilakukan commit ke database."""
    if not file.filename.endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="Format file harus .xlsx")

    contents = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return {"total_rows": 0, "valid_count": 0, "error_count": 0, "rows": []}

    parsed_rows = []
    valid_count = 0
    error_count = 0

    for idx, r in enumerate(rows[1:], start=2):
        if not r or all(cell is None for cell in r):
            continue

        nik = str(r[1]).strip() if len(r) > 1 and r[1] is not None else ""
        no_kk = str(r[2]).strip() if len(r) > 2 and r[2] is not None else ""
        nis = str(r[3]).strip() if len(r) > 3 and r[3] is not None else ""
        nisn = str(r[4]).strip() if len(r) > 4 and r[4] is not None else ""
        nama = str(r[5]).strip() if len(r) > 5 and r[5] is not None else ""
        jk = str(r[6]).strip().upper() if len(r) > 6 and r[6] is not None else ""
        tempat_lahir = str(r[7]).strip() if len(r) > 7 and r[7] is not None else ""
        tanggal_lahir = str(r[8]).strip() if len(r) > 8 and r[8] is not None else ""
        nama_ibu = str(r[9]).strip() if len(r) > 9 and r[9] is not None else ""
        nama_ayah = str(r[10]).strip() if len(r) > 10 and r[10] is not None else ""
        jalan = str(r[11]).strip() if len(r) > 11 and r[11] is not None else ""
        desa = str(r[12]).strip() if len(r) > 12 and r[12] is not None else ""
        kecamatan = str(r[13]).strip() if len(r) > 13 and r[13] is not None else ""
        kabupaten = str(r[14]).strip() if len(r) > 14 and r[14] is not None else ""
        provinsi = str(r[15]).strip() if len(r) > 15 and r[15] is not None else ""

        errors = []
        if not nama:
            errors.append("Nama lengkap wajib diisi")
        if not nis:
            errors.append("NIS wajib diisi")
        if jk not in ["L", "P"]:
            errors.append("JK harus 'L' atau 'P'")

        # Periksa duplikasi NIS di database
        if nis:
            existing_nis = db.exec(select(Siswa).where(Siswa.nis == nis, Siswa.deleted == 0)).first()
            if existing_nis:
                errors.append(f"NIS '{nis}' sudah terdaftar ({existing_nis.nama_lengkap})")

        is_valid = len(errors) == 0
        if is_valid:
            valid_count += 1
        else:
            error_count += 1

        parsed_rows.append({
            "row_index": idx,
            "is_valid": is_valid,
            "errors": errors,
            "data": {
                "nik": nik,
                "no_kk": no_kk,
                "nis": nis,
                "nisn": nisn,
                "nama_lengkap": nama,
                "jk": jk,
                "tempat_lahir": tempat_lahir,
                "tanggal_lahir": tanggal_lahir,
                "nama_ibu": nama_ibu,
                "nama_ayah": nama_ayah,
                "jalan": jalan,
                "desa_nama": desa,
                "kecamatan_nama": kecamatan,
                "kabupaten_nama": kabupaten,
                "provinsi_nama": provinsi,
            }
        })

    return {
        "total_rows": len(parsed_rows),
        "valid_count": valid_count,
        "error_count": error_count,
        "rows": parsed_rows,
    }
