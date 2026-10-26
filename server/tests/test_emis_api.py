import io
import openpyxl
from fastapi.testclient import TestClient
from sqlmodel import Session

def test_emis_template_and_preview(client: TestClient, session: Session):
    # 1. Setup Admin & Login
    client.post("/api/v1/auth/setup-admin", json={
        "username": "admin_emis",
        "pin": "123456",
        "nama_madrasah": "MTS ROUDHOTUT THOWALIB"
    })
    
    login_resp = client.post("/api/v1/auth/login", json={
        "username": "admin_emis",
        "pin": "123456",
        "device_id": "operator-pc",
        "nama_perangkat": "PC Operator"
    })
    assert login_resp.status_code == 200
    token = login_resp.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    # 2. Download template
    tmpl_resp = client.get("/api/v1/emis/template", headers=headers)
    assert tmpl_resp.status_code == 200
    assert "spreadsheetml" in tmpl_resp.headers["content-type"]

    # Verify template can be loaded by openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(tmpl_resp.content))
    assert "Format EMIS Siswa" in wb.sheetnames
    ws = wb["Format EMIS Siswa"]
    assert ws.cell(row=1, column=1).value == "NO"
    assert ws.cell(row=1, column=4).value == "NIS"

    # 3. Create a test excel upload to preview
    test_wb = openpyxl.Workbook()
    test_ws = test_wb.active
    test_ws.append(["NO", "NIK", "NO KK", "NIS", "NISN", "NAMA LENGKAP", "JK", "TEMPAT LAHIR", "TANGGAL LAHIR (YYYY-MM-DD)", "NAMA IBU", "NAMA AYAH", "ALAMAT JALAN", "DESA/KELURAHAN", "KECAMATAN", "KABUPATEN/KOTA", "PROVINSI"])
    test_ws.append([1, "3301011205120001", "3301010101100002", "2001", "0081234599", "Muhammad Ali", "L", "Semarang", "2012-01-15", "Siti", "Hasan", "Jl. Kauman", "Kauman", "Semarang", "Semarang", "Jateng"])
    # Invalid row: JK is 'X'
    test_ws.append([2, "", "", "2002", "", "Zahra", "X", "", "", "", "", "", "", "", "", ""])

    upload_bytes = io.BytesIO()
    test_wb.save(upload_bytes)
    upload_bytes.seek(0)

    preview_resp = client.post(
        "/api/v1/emis/import-preview",
        headers=headers,
        files={"file": ("siswa_emis.xlsx", upload_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    assert preview_resp.status_code == 200
    data = preview_resp.json()
    assert data["total_rows"] == 2
    assert data["valid_count"] == 1
    assert data["error_count"] == 1
    assert data["rows"][0]["is_valid"] is True
    assert data["rows"][1]["is_valid"] is False
    assert any("JK" in err for err in data["rows"][1]["errors"])
